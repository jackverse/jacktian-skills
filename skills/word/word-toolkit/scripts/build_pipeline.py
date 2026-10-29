# -*- coding: utf-8 -*-
"""一键流水线：build_docx.py 构建 + 本模板专属后处理（标题 numPr、TOC 域归位、分节符恢复）。

针对《04自贸区张江片区产业数字底座项目项目总结报告》类模板的三个特性：
1. 标题编号靠段落级直接 numPr（numId=3），样式上的 numPr 是"第%1章"历史遗留；
2. "目录"标题段之前藏着空的 toc 样式段，build 的新 TOC 域会插到"目录"之前，需归位；
3. 封面节/目录节的分节符恰好挂在旧目录区段落上，删除旧目录会连带删掉分节符，
   需从模板收割 sectPr 回插（封面末尾空段 + 正文首个 Heading 之前）。

用法:
  python build_pipeline.py <template.docx> <content.md> <output.docx> -- <build_docx.py 的参数...>
"""
import io
import subprocess
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def post_fix(output, template, numid="3", nonum_regex=r"^附件"):
    from copy import deepcopy

    # --- 1. 标题直接 numPr ---
    d = Document(output)
    import re as _re
    fixed = nonum = 0
    for p in d.paragraphs:
        m = _re.match(r"Heading (\d)$", p.style.name if p.style else "")
        if not m or int(m.group(1)) > 6:
            continue
        level = int(m.group(1))
        ppr = p._p.get_or_add_pPr()
        old = ppr.find(qn("w:numPr"))
        if old is not None:
            ppr.remove(old)
        numpr = ppr.get_or_add_numPr()
        no_num = nonum_regex and _re.search(nonum_regex, p.text.strip())
        ilvl = numpr.find(qn("w:ilvl"))
        if ilvl is None:
            ilvl = numpr.makeelement(qn("w:ilvl"), {})
            numpr.append(ilvl)
        ilvl.set(qn("w:val"), str(level - 1))
        nid = numpr.find(qn("w:numId"))
        if nid is None:
            nid = numpr.makeelement(qn("w:numId"), {})
            numpr.append(nid)
        nid.set(qn("w:val"), "0" if no_num else numid)
        nonum += 1 if no_num else 0
        fixed += 0 if no_num else 1
    d.save(output)
    print("[post 1] numPr: %d headings, %d no-number" % (fixed, nonum))

    # --- 2. TOC 域移到"目录"标题之后 ---
    d = Document(output)
    toc_field = toc_title = None
    for ch in d.element.body.iter(qn("w:p")):
        if toc_field is None:
            for it in ch.iter(qn("w:instrText")):
                if "TOC" in (it.text or ""):
                    toc_field = ch
                    break
    for p in d.paragraphs:
        nm = p.style.name if p.style else ""
        if toc_title is None and p.text.strip() == "目录" and nm.lower().startswith("toc"):
            toc_title = p._p
    moved = False
    if toc_field is not None and toc_title is not None and toc_field.getprevious() is not toc_title:
        toc_title.addnext(toc_field)
        moved = True
    d.save(output)
    print("[post 2] TOC field moved after 目录 title:", moved)

    # --- 3. 分节符恢复（从模板收割 sectPr 回插）---
    tpl = Document(template)
    tchildren = list(tpl.element.body)
    tpl_sects = []
    for ch in tchildren:
        if ch.tag == qn("w:p"):
            sp = ch.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            if sp is not None:
                tpl_sects.append(deepcopy(sp))
    # 模板段落级 sectPr 依次为 封面节尾、目录节尾…最后一个是 body-last（不含）
    d = Document(output)
    body = d.element.body
    cur_paras = list(body)
    out_sects = sum(
        1 for ch in cur_paras
        if (ch.tag == qn("w:p") and ch.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None)
        or ch.tag == qn("w:sectPr")
    )
    need = len(tpl_sects) + 1 - out_sects  # 目标：1 个 body-last + N 个段落级
    if need > 0:
        id2name = {s.style_id: s.name for s in d.styles if s.style_id is not None}
        toc_title = first_h = None
        for ch in cur_paras:
            if ch.tag != qn("w:p"):
                continue
            txt = "".join(t.text or "" for t in ch.iter(qn("w:t"))).strip()
            sid = ch.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            nm = id2name.get(sid.get(qn("w:val"))) if sid is not None else "Normal"
            if toc_title is None and txt == "目录" and (nm or "").lower().startswith("toc"):
                toc_title = ch
            if first_h is None and (nm or "").startswith("Heading"):
                first_h = ch
        # 封面节尾：挂到目录标题前一个段落
        cover_end = toc_title.getprevious()
        ppr = cover_end.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            cover_end.insert(0, ppr)
        ppr.append(deepcopy(tpl_sects[0]))
        need -= 1
        # 其余（目录节尾）：插在正文首个 Heading 之前的空段
        while need > 0:
            br = OxmlElement("w:p")
            ppr2 = OxmlElement("w:pPr")
            ppr2.append(deepcopy(tpl_sects[1]))
            br.append(ppr2)
            first_h.addprevious(br)
            need -= 1
        d.save(output)
    d2 = Document(output)
    print("[post 3] sections restored:", len(d2.sections))


def main():
    if len(sys.argv) < 6 or sys.argv[4] != "--":
        print(__doc__)
        sys.exit(1)
    template, content_md, output = sys.argv[1], sys.argv[2], sys.argv[3]
    args = sys.argv[5:]
    build = __file__.replace("build_pipeline.py", "") if False else None
    import os
    build_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "build_docx.py")
    if not os.path.exists(build_script):
        build_script = r"C:\Users\33380\.zcode\skills\word-toolkit\scripts\build_docx.py"
    print("[build] build_docx.py ...")
    r = subprocess.run([sys.executable, build_script, template, content_md, output] + args)
    if r.returncode != 0:
        sys.exit(r.returncode)
    post_fix(output, template)


if __name__ == "__main__":
    main()
