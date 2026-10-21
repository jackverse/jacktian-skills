# -*- coding: utf-8 -*-
"""分析 Word 模板/文档的版式事实，输出一份构建前必读的报告。

产出内容（全部 python-docx 实测，不靠肉眼）：
  - 分节：页面尺寸、页边距、页眉/页脚文字与 PAGE 域检测、linked_to_previous
  - Normal / Heading 1-6 样式：字体、字号、加粗、对齐、段前段后、行距
  - "目录"标题段落（样式可能挂在 toc/自定义样式上，决定 build_docx 的定位）
  - 封面段落（前 N 个非空段落的 run 级字体，含中文字体 eastAsia）
  - 表 0（通常是文档修订记录表）的行列结构与前几行内容
  - 段落/表格/图片/媒体计数
  - 可选 --outline：全文标题大纲（Heading 段落），用于写内容前的结构参考

用法:
  python analyze_docx.py <docx> [--outline] [--cover-paras 14] [--out report.md]
不指定 --out 时打印到 stdout。
"""
import argparse
import datetime
import io
import os
import sys
import zipfile

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ALIGN = {
    None: "继承(默认左)",
    WD_ALIGN_PARAGRAPH.LEFT: "LEFT",
    WD_ALIGN_PARAGRAPH.CENTER: "CENTER",
    WD_ALIGN_PARAGRAPH.RIGHT: "RIGHT",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "JUSTIFY",
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: "DISTRIBUTE",
}


def emu_cm(v):
    """EMU -> cm，保留 2 位。"""
    return None if v is None else round(v / 360000.0, 2)


def east_asia_of(rpr_holder):
    """从持有 rPr 的元素（run 或 style）取中文字体 eastAsia。"""
    rpr = rpr_holder.find(qn("w:rPr"))
    if rpr is None:
        return None
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        return None
    return rf.get(qn("w:eastAsia"))


def style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def run_info(run, para_style):
    """run 级字体信息，None 时沿样式链回退。返回 (name, eastAsia, size_pt, bold)"""
    name, ea = run.font.name, east_asia_of(run._element)
    size = run.font.size.pt if run.font.size else None
    bold = run.font.bold
    if name is None or ea is None or size is None or bold is None:
        for st in style_chain(para_style):
            try:
                rpr = st.element.find(qn("w:rPr"))
            except Exception:
                rpr = None
            if name is None:
                name = st.font.name
            if ea is None and rpr is not None:
                rf = rpr.find(qn("w:rFonts"))
                if rf is not None:
                    ea = rf.get(qn("w:eastAsia"))
            if size is None and st.font.size:
                size = st.font.size.pt
            if bold is None:
                bold = st.font.bold
            if name and ea and size and bold is not None:
                break
    return name, ea, size, bold


def footer_has_field(part):
    """页脚部件是否含域代码（PAGE/NUMPAGES 等）。"""
    try:
        xml = part._element.xml
    except Exception:
        return False
    return "PAGE" in xml or "fldChar" in xml or "fldSimple" in xml


def analyze(path, outline=False, cover_paras=14):
    doc = Document(path)
    out = [f"# 版式事实（python-docx 实测）：{os.path.basename(path)}"]
    out.append(f"- 生成时间：{datetime.datetime.now():%Y-%m-%d %H:%M}")

    # ---- 分节与页眉脚 ----
    out.append(f"\n## 分节（共 {len(doc.sections)} 节）")
    for i, s in enumerate(doc.sections):
        out.append(f"### 节 {i}（page {emu_cm(s.page_width)}x{emu_cm(s.page_height)} cm）")
        out.append(
            f"- 页边距：上 {emu_cm(s.top_margin)} 下 {emu_cm(s.bottom_margin)} "
            f"左 {emu_cm(s.left_margin)} 右 {emu_cm(s.right_margin)} cm；"
            f"装订线 {emu_cm(s.gutter)}；页眉距边界 {emu_cm(s.header_distance)}；"
            f"页脚距边界 {emu_cm(s.footer_distance)}"
        )
        for tag, part in (
            ("页眉", s.header),
            ("页脚", s.footer),
            ("首页页眉", s.first_page_header),
            ("首页页脚", s.first_page_footer),
        ):
            try:
                txts = [f"{p.text.strip()}({ALIGN.get(p.alignment, p.alignment)})"
                        for p in part.paragraphs if p.text.strip()]
                linked = part.is_linked_to_previous
                has_fld = footer_has_field(part)
                desc = " | ".join(txts) if txts else "（空）"
                if has_fld:
                    desc += f"  [含域: PAGE 等]"
                out.append(f"- {tag}: {desc} [linked_to_previous={linked}]")
            except Exception as e:
                out.append(f"- {tag}: 读取失败 {e}")

    # ---- 样式 ----
    st = doc.styles["Normal"]
    ea = east_asia_of(st.element)
    pf = st.paragraph_format
    out.append("\n## Normal 样式")
    out.append(f"- font.name={st.font.name}, eastAsia={ea}, "
               f"size={st.font.size.pt if st.font.size else None}pt, bold={st.font.bold}")
    out.append(f"- 行距={pf.line_spacing}, "
               f"段前={emu_cm(pf.space_before) if pf.space_before else None}cm "
               f"段后={emu_cm(pf.space_after) if pf.space_after else None}cm, "
               f"首行缩进={emu_cm(pf.first_line_indent) if pf.first_line_indent else None}cm")

    out.append("\n## Heading 样式定义")
    for hname in [f"Heading {k}" for k in range(1, 7)]:
        try:
            h = doc.styles[hname]
        except KeyError:
            out.append(f"- {hname}: 【模板中不存在】")
            continue
        ea = east_asia_of(h.element)
        pf = h.paragraph_format
        ppr = h.element.find(qn("w:pPr"))
        pbb = ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None
        out.append(
            f"- {hname}: font={h.font.name}/eastAsia={ea}, "
            f"size={h.font.size.pt if h.font.size else None}pt, bold={h.font.bold}, "
            f"对齐={ALIGN.get(pf.alignment, pf.alignment)}, "
            f"段前={pf.space_before.pt if pf.space_before else None}pt "
            f"段后={pf.space_after.pt if pf.space_after else None}pt, "
            f"行距={pf.line_spacing}, 段前分页={pbb}"
        )

    # 表格样式盘点（build 需要知道有没有 Table Grid 可用）
    grid_like = [s.name for s in doc.styles
                 if s.type == 3 and ("grid" in s.name.lower() or "表格" in s.name or "table" in s.name.lower())]
    out.append(f"\n## 表格样式（含 grid/table 关键字）: {grid_like or '【无！build 会走兜底加边框】'}")

    # ---- 目录标题 ----
    out.append('\n## "目录"标题段落')
    found = 0
    for i, p in enumerate(doc.paragraphs[:120]):
        t = p.text.strip()
        if t.replace(" ", "") == "目录" or t.replace(" ", "").startswith("目录"):
            rn = p.runs[0] if p.runs else None
            info = run_info(rn, p.style) if rn else (None,) * 4
            out.append(f'- para[{i}] text="{t}" style=<{p.style.name}> '
                       f"对齐={ALIGN.get(p.alignment, p.alignment)} "
                       f"首run: font={info[0]}/eastAsia={info[1]}/{info[2]}pt bold={info[3]}")
            found += 1
            if found >= 2:
                break
    if not found:
        out.append("- 未在前 120 段找到“目录”段落（该模板可能无目录页，或目录藏在域里——"
                   "用 build 后 word2pdf 更新域确认）")

    # ---- 封面段落 ----
    out.append(f"\n## 封面段落（前 {cover_paras} 个非空段落）")
    cnt = 0
    for i, p in enumerate(doc.paragraphs[:60]):
        t = p.text.strip()
        if not t and not p._p.findall(".//" + qn("w:drawing")):
            continue
        rn = p.runs[0] if p.runs else None
        if rn is not None:
            info = run_info(rn, p.style)
            fi = f"font={info[0]}/ea={info[1]}/{info[2]}pt/bold={info[3]}"
        else:
            fi = "（无 run，仅样式）"
        out.append(f'- para[{i}] <{p.style.name}> 对齐={ALIGN.get(p.alignment, p.alignment)} '
                   f'"{t[:40]}" {fi}')
        cnt += 1
        if cnt >= cover_paras:
            break

    # ---- 表 0（一般是修订记录表）----
    out.append("\n## 表 0 结构（通常是文档修订记录表）")
    if doc.tables:
        tbl = doc.tables[0]
        nr, nc = len(tbl.rows), len(tbl.columns)
        out.append(f"- 共 {nr} 行 x {nc} 列；前 5 行：")
        out.append("```")
        for r in tbl.rows[:5]:
            out.append("| " + " | ".join(c.text.strip()[:24] for c in r.cells) + " |")
        out.append("```")
    else:
        out.append("- 无表格")

    # ---- 其他计数 ----
    out.append("\n## 其他")
    out.append(f"- 段落总数 {len(doc.paragraphs)}，表格总数 {len(doc.tables)}，节总数 {len(doc.sections)}")
    body = doc.element.body
    out.append(f"- 图片：wp:inline {len(body.findall('.//' + qn('wp:inline')))} 个 / "
               f"wp:anchor {len(body.findall('.//' + qn('wp:anchor')))} 个")
    try:
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith("word/media/")]
            flds = sum(1 for n in z.namelist()
                       if n == "word/document.xml" and b"fldSimple" in z.read(n))
        out.append(f"- word/media/ 文件 {len(media)} 个；document.xml 含 fldSimple(隐藏域，注意“幽灵目录”)："
                   f"{'是' if flds else '否'}")
    except Exception as e:
        out.append(f"- zip 检查失败 {e}")

    # ---- 可选大纲 ----
    if outline:
        out.append("\n## 标题大纲（Heading 段落）")
        for i, p in enumerate(doc.paragraphs):
            sn = p.style.name if p.style else ""
            if sn.startswith("Heading"):
                lv = sn.split()[-1]
                out.append(f"{'  ' * (int(lv) - 1)}- [{i}] ({sn}) {p.text.strip()[:60]}")

    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--outline", action="store_true", help="追加全文标题大纲")
    ap.add_argument("--cover-paras", type=int, default=14)
    ap.add_argument("--out", help="写入文件而非 stdout")
    args = ap.parse_args()

    report = analyze(args.docx, outline=args.outline, cover_paras=args.cover_paras)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(report + "\n")
        print("written:", args.out, os.path.getsize(args.out), "bytes")
    else:
        print(report)


if __name__ == "__main__":
    main()
