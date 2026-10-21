# -*- coding: utf-8 -*-
"""
把 Markdown 内容注入到 Word 模板副本中，生成与模板格式完全一致的新文档
（保留模板的封面、修订记录表、样式、页眉脚、分节；替换正文与目录）。

适用场景：已有固定格式的公司/行业 Word 模板，需要按新项目内容批量产出
验收文档（详细设计说明书、自测报告、用户操作手册、试运行方案等）。

用法:
  python build_docx.py <template.docx> <content.md> <output.docx> [--mode auto|manual]
    [--replace 旧文本=新文本 ...]          # 全文替换（封面项目名、日期等）
    [--rev-date YYYY.MM.DD] [--rev-ver V1.0] [--rev-desc 新建] [--rev-author 编写人]
                                           # 写入修订记录表第一行

典型调用（把模板 A 项目换成 B 项目）:
  python build_docx.py 模板.docx content.md 输出.docx --mode manual \
    --replace "某旧项目名称=某新项目名称" --replace "二〇二五年五月=二〇二六年十月" \
    --rev-date 2026.10.14 --rev-ver V1.0 --rev-desc 新建 --rev-author 项目组

Markdown 约定:
  #/##/###/####/#####/######  → Heading 1..6
  普通段落 → Normal 样式；支持 **加粗**
  | a | b | 管道表格（首行为表头，自动表头重复/行不拆分/表头底纹）
  ::IMG::说明文字 → 截图占位框（居中灰字单格边框表 + 居中题注）
  !PAGEBREAK → 分页
  行首编号会被剥离；manual 模式下 3 级及以下标题按上下文重算字面编号
"""
import argparse
import datetime
import io
import re
import shutil
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


def el(tag):
    return OxmlElement(tag)


def replace_in_paragraph(p, mapping):
    text = p.text
    if not text:
        return
    new = text
    hit = False
    for old, repl in mapping.items():
        if old in new:
            new = new.replace(old, repl)
            hit = True
    if hit and p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""


def replace_everywhere(doc, mapping):
    def walk_paragraphs(paragraphs, tables):
        for p in paragraphs:
            replace_in_paragraph(p, mapping)
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    walk_paragraphs(cell.paragraphs, cell.tables)

    for sec in doc.sections:
        walk_paragraphs(sec.header.paragraphs, sec.header.tables)
        walk_paragraphs(sec.footer.paragraphs, sec.footer.tables)
    walk_paragraphs(doc.paragraphs, doc.tables)


def style_id_name_map(doc):
    return {s.style_id: s.name for s in doc.styles if s.style_id is not None}


def para_style_name(ch, id2name):
    if ch.tag != qn("w:p"):
        return None
    ps = ch.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    sid = ps.get(qn("w:val")) if ps is not None else None
    return id2name.get(sid, "Normal") if sid else "Normal"


def para_text(ch):
    return "".join(t.text or "" for t in ch.iter(qn("w:t"))).strip()


def locate(doc, id2name):
    """返回 (body_start_idx, first_toc_idx, toc_indexes, rev_tbl_element)"""
    body = doc.element.body
    children = list(body)
    date_idx = None
    toc_paras = []  # (idx, text)
    for i, ch in enumerate(children):
        if ch.tag != qn("w:p"):
            continue
        txt = para_text(ch)
        if date_idx is None and ("二〇二五年" in txt or "2025年" in txt):
            date_idx = i
        name = para_style_name(ch, id2name) or ""
        if name.lower().startswith("toc"):
            toc_paras.append((i, txt))
    entry_tocs = [(i, t) for i, t in toc_paras if t != "目录"]
    if entry_tocs:
        toc_last = max(i for i, _ in entry_tocs)
        search_from = toc_last + 1
    elif date_idx is not None:
        search_from = date_idx + 1
    else:
        search_from = 0
    body_start = None
    for i in range(search_from, len(children)):
        name = para_style_name(children[i], id2name) or ""
        if name.startswith("Heading"):
            body_start = i
            break
    if body_start is None:
        raise RuntimeError("找不到正文起始 Heading")
    # 修订记录表：文档修订记录 段之后第一个表
    rev_tbl = None
    rev_idx = None
    for i, ch in enumerate(children):
        if ch.tag == qn("w:p") and para_text(ch).startswith("文档修订记录"):
            rev_idx = i
            break
    if rev_idx is not None:
        for ch in children[rev_idx + 1:]:
            if ch.tag == qn("w:tbl"):
                rev_tbl = ch
                break
    first_toc = min(i for i, _ in entry_tocs) if entry_tocs else None
    return body_start, first_toc, [i for i, _ in entry_tocs], rev_tbl


def make_toc_field_paragraph(levels="1-3"):
    p = el("w:p")
    ppr = el("w:pPr")
    p.append(ppr)
    r = el("w:r")
    f = el("w:fldChar")
    f.set(qn("w:fldCharType"), "begin")
    f.set(qn("w:dirty"), "true")
    r.append(f)
    p.append(r)
    r = el("w:r")
    it = el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "%s" \\h \\z \\u ' % levels
    r.append(it)
    p.append(r)
    r = el("w:r")
    f = el("w:fldChar")
    f.set(qn("w:fldCharType"), "separate")
    r.append(f)
    p.append(r)
    r = el("w:r")
    t = el("w:t")
    t.text = "（目录内容将在打开文档时自动更新；如未刷新，请全选后按 F9 更新域。）"
    r.append(t)
    p.append(r)
    r = el("w:r")
    f = el("w:fldChar")
    f.set(qn("w:fldCharType"), "end")
    r.append(f)
    p.append(r)
    return p


def set_update_fields(doc):
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = el("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.insert(0, uf)


HEAD_RE = re.compile(r"^(#{1,6})\s+(.*)$")
NUM_PREFIX_RE = re.compile(r"^\s*\d+(?:\.\d+)*\.?\s+")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def parse_md(md_text):
    blocks = []
    lines = md_text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        m = HEAD_RE.match(s)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s == "!PAGEBREAK":
            blocks.append(("pagebreak",))
            i += 1
            continue
        if s.startswith("::IMG::"):
            blocks.append(("img", s[len("::IMG::"):].strip()))
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            rows = [[c.strip() for c in s.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if s.startswith("- "):
            blocks.append(("p", "● " + s[2:].strip()))
            i += 1
            continue
        blocks.append(("p", s))
        i += 1
    return blocks


def apply_rpr(run, proto_rpr, force_bold=None):
    """把模板收割的 run 级 rPr 原型套到 run 上（模板正文/标题普遍带直接字体，
    仅靠样式注入会出现"字体变了"。force_bold: True=强制加粗，None=不动加粗。"""
    if proto_rpr is None:
        if force_bold:
            run.bold = True
        return
    from copy import deepcopy

    rpr = deepcopy(proto_rpr)
    for tag in ("w:b", "w:bCs"):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    if force_bold:
        b, bcs = el("w:b"), el("w:bCs")
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            rf.addnext(bcs)
            rf.addnext(b)
        else:
            rpr.insert(0, bcs)
            rpr.insert(0, b)
    old = run._r.find(qn("w:rPr"))
    if old is not None:
        run._r.remove(old)
    run._r.insert(0, rpr)


def add_runs(p, text, base_bold=False, proto_rpr=None):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            apply_rpr(r, proto_rpr, force_bold=True if base_bold else None)
        r = p.add_run(m.group(1))
        apply_rpr(r, proto_rpr, force_bold=True)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        apply_rpr(r, proto_rpr, force_bold=True if base_bold else None)


def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = el("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def new_table(doc, nrows, ncols, style):
    t = doc.add_table(rows=nrows, cols=ncols, style=style)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = el("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    return t


def add_table(doc, rows, table_style, anchor):
    ncols = max(len(r) for r in rows)
    t = new_table(doc, len(rows), ncols, table_style)
    if anchor is not None:
        anchor.addprevious(t._tbl)
    # 首行若只有 1 个单元格且列数>1，视为"功能描述"式合并表头行，跨列合并
    merge_first = len(rows[0]) == 1 and ncols > 1
    if merge_first:
        t.cell(0, 0).merge(t.cell(0, ncols - 1))
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            if merge_first and ri == 0:
                continue
            txt = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            add_runs(p, txt, base_bold=(ri == 0))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            if ri == 0:
                set_cell_shading(cell, "F2F2F2")
    if merge_first:
        cell0 = t.cell(0, 0)
        cell0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cell0.paragraphs[0].runs:
            r.bold = True
        set_cell_shading(cell0, "F2F2F2")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(el("w:tblHeader"))
    for row in t.rows:
        row._tr.get_or_add_trPr().append(el("w:cantSplit"))
    set_content_widths(t, rows, ncols, merge_first)


def set_content_widths(t, rows, ncols, merge_first):
    """按各列内容长度加权设置列宽（固定布局），避免均分导致的尴尬换行。"""
    data_rows = rows[1:] if merge_first else rows
    weights = []
    for ci in range(ncols):
        m = 0
        for row in data_rows:
            v = row[ci] if ci < len(row) else ""
            m = max(m, min(len(v), 20))
        weights.append(max(m, 4) ** 0.7)
    total = sum(weights)
    # A4 版心约 8312 twips（21cm - 2*3.17cm）
    TEXT_W = 8312
    widths = [max(int(TEXT_W * w / total), 1000) for w in weights]
    scale = TEXT_W / sum(widths)
    widths = [int(w * scale) for w in widths]
    tblPr = t._tbl.tblPr
    layout = el("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        for ci, gc in enumerate(cols):
            if ci < len(widths):
                gc.set(qn("w:w"), str(widths[ci]))
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            if ci < len(widths):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = el("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(widths[ci]))
                tcW.set(qn("w:type"), "dxa")


def add_img_placeholder(doc, caption, anchor, table_style="Table Grid"):
    t = new_table(doc, 1, 1, table_style)
    if anchor is not None:
        anchor.addprevious(t._tbl)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trH = el("w:trHeight")
    trH.set(qn("w:val"), "4000")
    trH.set(qn("w:hRule"), "atLeast")
    trPr.append(trH)
    cell = t.cell(0, 0)
    set_cell_shading(cell, "FAFAFA")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("【截图占位：%s】" % caption)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    cap = doc.add_paragraph()
    if anchor is not None:
        anchor.addprevious(cap._p)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(cap, caption)


def pick_table_style(doc):
    counts = {}
    for t in doc.tables:
        try:
            n = t.style.name if t.style is not None else None
        except Exception:
            n = None
        if n:
            counts[n] = counts.get(n, 0) + 1
    if counts:
        return max(counts.items(), key=lambda kv: kv[1])[0]
    # 正文表已删空时，从样式表中选一个表格样式
    prefer = ["Table Grid", "网格型", "Normal Table"]
    table_styles = []
    for s in doc.styles:
        try:
            if s.type is not None and "TABLE" in str(s.type):
                table_styles.append(s.name)
        except Exception:
            continue
    for p in prefer:
        if p in table_styles:
            return p
    return table_styles[0] if table_styles else None


def harvest_prototypes(children, body_start, id2name):
    """从模板旧正文收割直接格式原型（在删除旧正文前调用）：
    - body_rpr：正文 run 的 rPr（宋体/Times New Roman 等直接字体）
    - ind：正文段的首行缩进 w:ind（如 firstLine=480 即首行缩进 2 字符）
    - head[lvl]：各级标题 run 的 rPr（黑体+字号等）
    模板的正文/标题大多不靠样式而靠直接格式定字体，仅挂样式注入会出现
    "段首不空两格、标题字体变了"。取最长且非加粗的 Normal 段做正文原型。"""
    from copy import deepcopy

    proto = {"body_rpr": None, "ind": None, "head": {}}
    best = 0
    for ch in children[body_start:]:
        if ch.tag != qn("w:p"):
            continue
        ppr = ch.find(qn("w:pPr"))
        sid = None
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            if ps is not None:
                sid = ps.get(qn("w:val"))
        name = id2name.get(sid, "Normal") if sid else "Normal"
        text = "".join(t.text or "" for t in ch.iter(qn("w:t"))).strip()
        if not text:
            continue
        runs = ch.findall(qn("w:r"))
        run0 = runs[0] if runs else None
        rpr0 = run0.find(qn("w:rPr")) if run0 is not None else None
        if name == "Normal":
            if len(text) > best and rpr0 is not None and rpr0.find(qn("w:b")) is None:
                best = len(text)
                proto["body_rpr"] = deepcopy(rpr0)
                ind = ppr.find(qn("w:ind")) if ppr is not None else None
                if ind is not None and (ind.get(qn("w:firstLine")) or ind.get(qn("w:firstLineChars"))):
                    proto["ind"] = deepcopy(ind)
        else:
            m = re.match(r"Heading (\d)$", name)
            if m:
                lvl = int(m.group(1))
                if lvl not in proto["head"] and rpr0 is not None:
                    proto["head"][lvl] = deepcopy(rpr0)
    return proto


def apply_body_indent(p, proto_ind):
    """给正文段套首行缩进原型（居中段/题注不套）。"""
    if proto_ind is None:
        return
    from copy import deepcopy

    ppr = p._p.get_or_add_pPr()
    old = ppr.find(qn("w:ind"))
    if old is not None:
        ppr.remove(old)
    ppr.append(deepcopy(proto_ind))


def heading_text(text, level, mode, counters):
    clean = NUM_PREFIX_RE.sub("", text).strip()
    if mode != "manual":
        return clean
    # manual 模式：1/2 级由样式自动编号，3 级以下按上下文补字面编号（与样式编号同序）
    counters[level] = counters.get(level, 0) + 1
    for l in list(counters):
        if l > level:
            del counters[l]
    if level <= 2:
        return clean
    prefix = ".".join(str(counters.get(l, 1)) for l in range(1, level + 1))
    return "%s %s" % (prefix, clean)


def strip_unused_media(path):
    """删除 docx 包中未被任何 part 引用的媒体文件，减小体积（模板旧截图已随旧正文删除）。"""
    import zipfile
    import tempfile
    import os as _os
    from lxml import etree

    R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    tmp = tempfile.mktemp(suffix=".docx")
    shutil.copyfile(path, tmp)
    zin = zipfile.ZipFile(tmp, "r")
    names = set(zin.namelist())

    # part 名 -> 该 part 实际引用的 media 集合（zip 内路径）
    used = set()
    for rf in [n for n in names if n.endswith(".rels")]:
        part = rf.replace("_rels/", "").replace(".rels", "")
        if part not in names:
            continue
        try:
            rel_root = etree.fromstring(zin.read(rf))
            part_root = etree.fromstring(zin.read(part))
        except Exception:
            continue
        base = _os.path.dirname(part)  # word
        rid2file = {}
        for rel in rel_root:
            t = rel.get("Target") or ""
            mode = rel.get("TargetMode") or "Internal"
            if mode != "Internal" or "media" not in t:
                continue
            # 解析相对路径（含 ../）
            combined = _os.path.normpath(_os.path.join(base, t)).replace("\\", "/")
            rid2file[rel.get("Id")] = combined
        for el in part_root.iter():
            for attr in ("{%s}embed" % R_NS, "{%s}id" % R_NS, "{%s}link" % R_NS, "{%s}pict" % R_NS):
                rid = el.get(attr)
                if rid and rid in rid2file:
                    used.add(rid2file[rid])

    all_media = [n for n in names if n.startswith("word/media/")]
    drop = set(m for m in all_media if m not in used)
    # rels 中删除指向被删 media 的 Relationship（避免悬空）
    fixed_rels = {}
    for rf in [n for n in names if n.endswith(".rels")]:
        try:
            root = etree.fromstring(zin.read(rf))
        except Exception:
            continue
        changed = False
        for rel in list(root):
            t = rel.get("Target") or ""
            if "media" in t and (rel.get("TargetMode") or "Internal") == "Internal":
                combined = _os.path.normpath(
                    _os.path.join(_os.path.dirname(rf.replace("_rels/", "")), t)
                ).replace("\\", "/")
                if combined in drop:
                    root.remove(rel)
                    changed = True
        if changed:
            fixed_rels[rf] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    zout = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    dropped = 0
    for item in zin.infolist():
        if item.filename in drop:
            dropped += 1
            continue
        data = fixed_rels.get(item.filename) or zin.read(item.filename)
        zout.writestr(item.filename, data)
    zout.close()
    zin.close()
    _os.remove(tmp)
    return dropped, len(all_media)


def ensure_footer_page_number(doc):
    """确保最后一个（正文）节的页脚显式包含居中 PAGE 页码域。
    模板正文的页码页脚常挂在随旧正文一起删除的分节符上，这里兜底重建。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    sec = doc.sections[-1]
    footer = sec.footer
    if not footer.is_linked_to_previous:
        xml = footer._element.xml if footer._element is not None else ""
        if "PAGE" in xml:
            return False
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r" PAGE   \* MERGEFORMAT ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("content_md")
    ap.add_argument("output")
    ap.add_argument("--mode", choices=["auto", "manual"], default="auto")
    ap.add_argument("--replace", action="append", default=[],
                    help="全文替换，格式 旧文本=新文本，可多次出现")
    ap.add_argument("--rev-date", default=datetime.date.today().strftime("%Y.%m.%d"))
    ap.add_argument("--rev-ver", default="V1.0")
    ap.add_argument("--rev-desc", default="新建")
    ap.add_argument("--rev-author", default="")
    args = ap.parse_args()

    shutil.copyfile(args.template, args.output)
    doc = Document(args.output)
    id2name = style_id_name_map(doc)

    mapping = {}
    for kv in args.replace:
        old, _, new = kv.partition("=")
        if old:
            mapping[old] = new

    body_start, first_toc, toc_idxs, rev_tbl = locate(doc, id2name)
    body = doc.element.body
    children = list(body)
    # 删旧正文前先收割直接格式原型（首行缩进/正文字体/标题字体）
    proto = harvest_prototypes(children, body_start, id2name)

    # 先插入新 TOC 域（此时旧条目尚未删除，锚点仍挂在 body 上）
    if first_toc is not None:
        toc_p = make_toc_field_paragraph()
        children[first_toc].addprevious(toc_p)
        # Heading 1 样式无"段前分页"时，目录后补一个分页符，避免正文与目录同页
        try:
            h1 = doc.styles["Heading 1"]
            h1_ppr = h1.element.find(qn("w:pPr"))
            h1_breaks = h1_ppr is not None and h1_ppr.find(qn("w:pageBreakBefore")) is not None
        except KeyError:
            h1_breaks = False
        if not h1_breaks:
            pb = el("w:p")
            ppr = el("w:pPr")
            pb.append(ppr)
            r = el("w:r")
            f = el("w:fldChar")
            f.set(qn("w:fldCharType"), "begin")
            r.append(f)
            it = el("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " PAGE "
            r.append(it)
            f2 = el("w:fldChar")
            f2.set(qn("w:fldCharType"), "end")
            r.append(f2)
            # 用简单分页 run 替换上面的域 runs（走 WD_BREAK）
            for child in list(r):
                r.remove(child)
            from docx.enum.text import WD_BREAK as _WB

            br = el("w:br")
            br.set(qn("w:type"), "page")
            r.append(br)
            pb.append(r)
            toc_p.addnext(pb)

    # 删除旧正文（保留末尾 sectPr）
    for ch in children[body_start:]:
        if ch.tag == qn("w:sectPr") and ch is children[-1]:
            continue
        body.remove(ch)
    # 删除旧 TOC 条目段
    for i in toc_idxs:
        ch = children[i]
        if ch is not toc_p and ch.getparent() is body:
            body.remove(ch)

    # 追加新正文
    with open(args.content_md, encoding="utf-8") as f:
        blocks = parse_md(f.read())
    sectPr = body.find(qn("w:sectPr"))
    anchor = sectPr  # 直接子级 sectPr（正文节），找不到则为 None
    table_style = pick_table_style(doc)
    counters = {}
    for blk in blocks:
        if blk[0] == "h":
            _, level, text = blk
            p = doc.add_paragraph(style="Heading %d" % min(level, 9))
            if anchor is not None:
                anchor.addprevious(p._p)
            add_runs(p, heading_text(text, level, args.mode, counters),
                     proto_rpr=proto["head"].get(min(level, 6)))
        elif blk[0] == "p":
            p = doc.add_paragraph()
            if anchor is not None:
                anchor.addprevious(p._p)
            add_runs(p, blk[1], proto_rpr=proto["body_rpr"])
            apply_body_indent(p, proto["ind"])
        elif blk[0] == "table":
            add_table(doc, blk[1], table_style, anchor)
        elif blk[0] == "img":
            add_img_placeholder(doc, blk[1], anchor, table_style)
        elif blk[0] == "pagebreak":
            p = doc.add_paragraph()
            if anchor is not None:
                anchor.addprevious(p._p)
            p.add_run().add_break(WD_BREAK.PAGE)

    # 修订记录表：表结构为 项目名称/编制单位 两行 + 版本表头 + 记录行/备用行
    # 清空模板历史记录，写入本次修订（保留表格行数不膨胀）
    if rev_tbl is not None:
        from docx.table import Table

        tbl = Table(rev_tbl, doc)
        vals = [args.rev_ver, args.rev_date, args.rev_desc, args.rev_author, "", ""]
        data_rows = tbl.rows[3:]
        target = None
        for ri, row in enumerate(data_rows):
            if all(not c.text.strip() for c in row.cells):
                if target is None:
                    target = row
                    for ci, cell in enumerate(row.cells):
                        cell.paragraphs[0].text = vals[ci] if ci < len(vals) else ""
                continue
            for ci, cell in enumerate(row.cells):
                cell.paragraphs[0].text = ""

    replace_everywhere(doc, mapping)
    set_update_fields(doc)
    try:
        footer_fixed = ensure_footer_page_number(doc)
    except Exception as e:
        footer_fixed = False
        print("footer warn:", e)
    doc.save(args.output)
    try:
        dropped, total = strip_unused_media(args.output)
    except Exception as e:
        dropped, total = -1, -1
        print("strip media warn:", e)
    print("saved:", args.output, "| body_start=", body_start, "| toc paras=", len(toc_idxs),
          "| rev_row=", rev_tbl is not None, "| tables=", len(doc.tables),
          "| paras=", len(doc.paragraphs), "| media dropped=", dropped, "/", total,
          "| footer_fixed=", footer_fixed,
          "| proto: body_rpr=", proto["body_rpr"] is not None, "ind=", proto["ind"] is not None,
          "head=", sorted(proto["head"]))


if __name__ == "__main__":
    main()
